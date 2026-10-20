"""Import the tool's vocabulary CSV into the database.

The CSV schema (see AI_Reader_Document/STANDARD.md):
    word, [traditional], level, pos, ipa, definition_en, meaning_vi, example, count

Language is auto-detected: presence of a `traditional` column => Chinese ("zh"),
otherwise English ("en"). It can be overridden explicitly.
"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, BinaryIO, TextIO

from sqlmodel import Session, delete, select

from .models import Source, Vocabulary


def detect_language(headers: list[str]) -> str:
    return "zh" if "traditional" in headers else "en"


def _get(row: dict[str, str], key: str) -> str:
    return (row.get(key) or "").strip()


def import_rows(
    session: Session,
    headers: list[str],
    raw_rows: list[list[str]],
    source_name: str,
    language: str | None = None,
    level: str | None = None,
) -> dict[str, Any]:
    """Insert rows for one source. Re-importing the same name replaces it."""
    language = language or detect_language(headers)
    header_index = {name: i for i, name in enumerate(headers)}

    def field(row: list[str], name: str) -> str:
        i = header_index.get(name)
        return (row[i] or "").strip() if i is not None and i < len(row) else ""

    # Idempotent: drop a previous import with the same name (and its words).
    existing = session.exec(select(Source).where(Source.name == source_name)).all()
    for src in existing:
        session.exec(delete(Vocabulary).where(Vocabulary.source_id == src.id))
        session.delete(src)
    session.commit()

    source = Source(name=source_name, language=language, level=level or "")
    session.add(source)
    session.commit()
    session.refresh(source)

    seen: set[str] = set()
    imported = 0
    skipped = 0
    levels: set[str] = set()
    for row in raw_rows:
        word = field(row, "word")
        if not word or word in seen:
            skipped += 1
            continue
        seen.add(word)
        row_level = field(row, "level") or (level or "")
        levels.add(row_level)
        session.add(
            Vocabulary(
                word=word,
                traditional=field(row, "traditional"),
                language=language,
                level=row_level,
                pos=field(row, "pos"),
                ipa=field(row, "ipa"),
                definition_en=field(row, "definition_en"),
                meaning_vi=field(row, "meaning_vi"),
                example=field(row, "example"),
                source_id=source.id,
            )
        )
        imported += 1

    if level:
        source.level = level
    elif len(levels) == 1:
        source.level = next(iter(levels))
    source.count = imported
    session.add(source)
    session.commit()
    session.refresh(source)

    return {
        "source_id": source.id,
        "source": source.name,
        "language": language,
        "level": source.level,
        "imported": imported,
        "skipped": skipped,
    }


def import_csv_file(
    session: Session,
    path: str | Path,
    source_name: str | None = None,
    language: str | None = None,
    level: str | None = None,
) -> dict[str, Any]:
    path = Path(path)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        rows = list(reader)
    if not rows:
        raise ValueError(f"Empty CSV: {path}")
    headers = [h.strip() for h in rows[0]]
    body = [r + [""] * (len(headers) - len(r)) for r in rows[1:]]
    return import_rows(
        session, headers, body, source_name or path.name,
        language=language, level=level,
    )


def import_xlsx_file(
    session: Session,
    path: str | Path,
    source_name: str | None = None,
    language: str | None = None,
    level: str | None = None,
    sheet: str | int | None = None,
) -> dict[str, Any]:
    """Import vocabulary from an Excel file (first sheet by default).

    Requires ``openpyxl``. Extra columns beyond the CSV schema are ignored.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "openpyxl is required for .xlsx import. Run: pip install openpyxl"
        ) from exc

    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            ws = wb[wb.sheetnames[0]]
        elif isinstance(sheet, int):
            ws = wb[wb.sheetnames[sheet]]
        else:
            ws = wb[sheet]
        rows = [
            ["" if c is None else str(c) for c in row]
            for row in ws.iter_rows(values_only=True)
        ]
    finally:
        wb.close()

    if not rows:
        raise ValueError(f"Empty XLSX: {path}")
    headers = [h.strip() for h in rows[0]]
    body = [
        r + [""] * (len(headers) - len(r))
        for r in rows[1:]
        if any(cell.strip() for cell in r)
    ]
    return import_rows(
        session, headers, body, source_name or path.stem,
        language=language, level=level,
    )


def import_file(
    session: Session,
    path: str | Path,
    source_name: str | None = None,
    language: str | None = None,
    level: str | None = None,
) -> dict[str, Any]:
    """Import ``.csv`` or ``.xlsx`` vocabulary into the database."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return import_csv_file(
            session, path, source_name=source_name, language=language, level=level
        )
    if suffix in (".xlsx", ".xlsm"):
        return import_xlsx_file(
            session, path, source_name=source_name, language=language, level=level
        )
    raise ValueError(f"Unsupported file type: {suffix} (use .csv or .xlsx)")


def import_csv_stream(
    session: Session,
    stream: BinaryIO | TextIO | bytes | str,
    source_name: str,
    language: str | None = None,
    level: str | None = None,
) -> dict[str, Any]:
    """Import from uploaded bytes/str or a file object (e.g. UploadFile.file)."""
    if isinstance(stream, (bytes, bytearray)):
        text = bytes(stream).decode("utf-8-sig")
    elif isinstance(stream, str):
        text = stream
    elif hasattr(stream, "read"):
        data = stream.read()
        text = data.decode("utf-8-sig") if isinstance(data, (bytes, bytearray)) else data
    else:  # pragma: no cover
        raise ValueError("Unsupported input for import")
    reader = csv.reader(text.splitlines())
    rows = list(reader)
    if not rows:
        raise ValueError("Empty upload")
    headers = [h.strip() for h in rows[0]]
    body = [r + [""] * (len(headers) - len(r)) for r in rows[1:]]
    return import_rows(
        session, headers, body, source_name, language=language, level=level
    )
